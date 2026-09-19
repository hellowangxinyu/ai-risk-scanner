"""客户维护：增删改查 + Excel/文本批量导入 + 模板下载。"""
import io
import urllib.parse

import openpyxl
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import db
from logic import validate_import_row
from scanner import list_customers_with_due

router = APIRouter(prefix="/api/customers", tags=["customers"])

TEMPLATE_HEADERS = ["名称", "类型", "统一社会信用代码", "联系人", "备注", "是否授信客户"]


class CustomerBody(BaseModel):
    name: str
    ctype: str = "公司"
    credit_code: str = ""
    contact: str = ""
    note: str = ""
    org: str = ""
    is_credit: bool = False


def _resolve_org(org: str = "") -> str:
    """机构名：显式指定优先，否则取系统配置 default_org（开源默认'本公司'）。"""
    org = (org or "").strip()
    return org or db.get_settings().get("default_org") or "本公司"


def _check_unique(conn, org: str, name: str, exclude_id: int = 0):
    row = conn.execute(
        "SELECT id FROM customers WHERE org=? AND name=? AND id!=?", (org, name, exclude_id)
    ).fetchone()
    if row:
        raise HTTPException(409, f"同机构下已存在同名客户：{name}")


@router.get("")
def list_customers(query: str = "", page: int = 1, page_size: int = 20):
    items = list_customers_with_due()
    if query.strip():
        q = query.strip()
        items = [
            c
            for c in items
            if q in c["name"] or q in c["contact"] or q in c["credit_code"]
        ]
    total = len(items)
    page = max(page, 1)
    page_size = max(min(page_size, 200), 1)
    start = (page - 1) * page_size
    due_count = sum(1 for c in items if c["due"])
    return {"total": total, "due_count": due_count, "items": items[start : start + page_size]}


@router.post("")
def create_customer(body: CustomerBody):
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "客户名称不能为空")
    if body.ctype not in ("公司", "个人"):
        raise HTTPException(400, "类型必须为 公司 或 个人")
    org = _resolve_org(body.org)
    conn = db.connect()
    try:
        _check_unique(conn, org, name)
        cur = conn.execute(
            "INSERT INTO customers(org, name, ctype, credit_code, contact, note, is_credit) VALUES(?,?,?,?,?,?,?)",
            (org, name, body.ctype, body.credit_code.strip(), body.contact.strip(), body.note.strip(), 1 if body.is_credit else 0),
        )
        conn.commit()
        return {"id": cur.lastrowid}
    finally:
        conn.close()


@router.put("/{cid}")
def update_customer(cid: int, body: CustomerBody):
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "客户名称不能为空")
    if body.ctype not in ("公司", "个人"):
        raise HTTPException(400, "类型必须为 公司 或 个人")
    conn = db.connect()
    try:
        if not conn.execute("SELECT id FROM customers WHERE id=?", (cid,)).fetchone():
            raise HTTPException(404, "客户不存在")
        org = _resolve_org(body.org)
        _check_unique(conn, org, name, exclude_id=cid)
        conn.execute(
            "UPDATE customers SET org=?, name=?, ctype=?, credit_code=?, contact=?, note=?, is_credit=? WHERE id=?",
            (org, name, body.ctype, body.credit_code.strip(), body.contact.strip(), body.note.strip(), 1 if body.is_credit else 0, cid),
        )
        conn.commit()
        return {"ok": True}
    finally:
        conn.close()


@router.delete("/{cid}")
def delete_customer(cid: int):
    conn = db.connect()
    try:
        conn.execute("DELETE FROM customers WHERE id=?", (cid,))
        conn.commit()
        return {"ok": True}
    finally:
        conn.close()


@router.get("/template")
def download_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "客户导入"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["示例建材有限公司", "公司", "91330100XXXXXXXXXX", "张三", "仅作格式示例，导入前请删除本行", "是"])
    ws.append(["李四", "个人", "", "", "个人客户可不填信用代码", "否"])
    ws.append(["示例供应商有限公司", "公司", "", "", "非授信客户按低频周期扫描", "否"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = urllib.parse.quote("客户导入模板.xlsx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


def _parse_rows(report: dict, rows_iter, has_header: bool):
    """逐行校验并写入，产出报告。文本导入 has_header=False，每行一个名称。"""
    from logic import parse_credit_flag

    org = _resolve_org()
    conn = db.connect()
    seen_in_file = set()
    try:
        for row_no, values in rows_iter:
            if has_header:
                ok, fields, reason = validate_import_row(
                    values.get("名称"), values.get("类型"), values.get("统一社会信用代码"),
                    values.get("联系人"), values.get("备注"),
                )
                name = (values.get("名称") or "").strip()
                if ok:
                    try:
                        fields["is_credit"] = 1 if parse_credit_flag(values.get("是否授信客户")) else 0
                    except ValueError as e:
                        ok, fields, reason = False, None, str(e)
            else:
                name = str(values or "").strip()
                ok, fields, reason = validate_import_row(name, "公司", "", "", "")
                if ok:
                    fields["is_credit"] = 0
            if not ok:
                report["failed"].append({"row": row_no, "name": name, "reason": reason})
                continue
            if fields["name"] in seen_in_file:
                report["skipped"].append({"name": fields["name"], "reason": "文件内重复"})
                continue
            dup = conn.execute(
                "SELECT id FROM customers WHERE org=? AND name=?", (org, fields["name"])
            ).fetchone()
            if dup:
                report["skipped"].append({"name": fields["name"], "reason": "已存在同名客户"})
                continue
            conn.execute(
                "INSERT INTO customers(org, name, ctype, credit_code, contact, note, is_credit) VALUES(?,?,?,?,?,?,?)",
                (org, fields["name"], fields["ctype"], fields["credit_code"], fields["contact"], fields["note"], fields["is_credit"]),
            )
            seen_in_file.add(fields["name"])
            report["success"] += 1
        conn.commit()
    finally:
        conn.close()


@router.post("/import")
async def import_excel(file: UploadFile = File(None)):
    if file is None:
        raise HTTPException(400, "请上传 Excel 文件")
    data = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(400, f"无法解析 Excel 文件：{e}")
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    headers = [str(h).strip() if h is not None else "" for h in (header or [])]
    if "名称" not in headers:
        raise HTTPException(400, "Excel 缺少「名称」列，请使用模板（可从客户管理页下载）")
    idx = {h: i for i, h in enumerate(headers)}
    report = {"success": 0, "skipped": [], "failed": []}

    def gen():
        for row_no, row in enumerate(rows, start=2):
            def val(col):
                i = idx.get(col)
                if i is None or i >= len(row):
                    return ""
                v = row[i]
                return "" if v is None else v

            yield row_no, {h: val(h) for h in TEMPLATE_HEADERS}

    _parse_rows(report, gen(), has_header=True)
    return report


@router.post("/import-text")
def import_text(text: str = Form(...)):
    report = {"success": 0, "skipped": [], "failed": []}

    def gen():
        for row_no, line in enumerate(text.splitlines(), start=1):
            if line.strip():
                yield row_no, line

    _parse_rows(report, gen(), has_header=False)
    return report
